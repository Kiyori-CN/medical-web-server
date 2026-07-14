from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook


HEADERS = [
    "题型",
    "题目",
    "考点还原",
    "答案",
    "答案解析",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
]


def _write_workbook(path: Path, rows: list[list[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "题目"
    worksheet.append(HEADERS)
    for row in rows:
        worksheet.append(row)
    workbook.save(path)
    workbook.close()


@pytest.fixture
def sample_source_tree(tmp_path: Path) -> Path:
    source_root = tmp_path / "data"
    _write_workbook(
        source_root / "执业医师" / "学科刷题" / "生理学" / "第一章.xlsx",
        [
            [
                "A1型题",
                "2020N1A 细胞膜运输方式属于",
                "膜转运基本概念",
                "A",
                "顺浓度梯度、不消耗能量。",
                "易化扩散",
                "主动转运",
                "胞吞",
                "胞吐",
                None,
            ],
            [
                "填空题",
                "维持内环境稳定的调节方式是______。",
                "稳态调节",
                "神经-体液调节",
                "机体通过多种调节方式维持稳态。",
                None,
                None,
                None,
                None,
                None,
            ],
            [
                "A3/A4型题",
                "2007N179A-临床 (临床类)女性,35岁,产后哺乳期,右乳红肿。查体体温36.8°C,乳房发硬无压痛,右腋下触及活动性淋巴结。初步诊断是",
                "病例串题",
                "A",
                "病例第一问。",
                "炎性乳癌",
                "急性乳腺炎",
                "乳腺囊肿",
                "乳腺纤维瘤",
                None,
            ],
            [
                "A3/A4型题",
                "2007N180A-临床 (临床类)女性,35岁,产后哺乳期,右乳红肿。查体体温36.8°C,乳房发硬无压痛,右腋下触及活动性淋巴结。明确诊断后首选治疗是",
                "病例串题",
                "A",
                "病例第二问。",
                "化学治疗",
                "单纯切开引流",
                "随访观察",
                "局部热敷",
                None,
            ],
        ],
    )
    _write_workbook(
        source_root / "执业医师" / "系统刷题" / "生理系统" / "重复题.xlsx",
        [
            [
                "单选题",
                "细胞膜运输方式属于",
                "膜转运基本概念",
                "B",
                "此行故意制造答案冲突，用于验证报告。",
                "A. 易化扩散",
                "B. 主动转运",
                "C. 胞吞",
                "D. 胞吐",
                None,
            ]
        ],
    )

    # Same display name as 第一章.xlsx but a different source format. Each
    # physical file remains an independent collection.
    docx_path = source_root / "执业医师" / "学科刷题" / "生理学" / "第一章.docx"
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    for text in (
        "1.[单选题]大气道狭窄时最典型的体征是",
        "A. 三凹征",
        "B. 颈静脉怒张",
        "C. 桶状胸",
        "D. 杵状指",
        "答案：A",
        "考点还原：吸气性呼吸困难可出现三凹征。",
        "解析：大气道狭窄主要表现为吸气性呼吸困难。",
    ):
        document.add_paragraph(text)
    document.save(docx_path)
    return source_root
