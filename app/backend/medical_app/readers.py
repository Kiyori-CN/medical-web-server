from __future__ import annotations

import hashlib
import re
from collections.abc import Iterator
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from .models import RawQuestion, SourceFile
from .normalize import display_text


SUPPORTED_EXTENSIONS = {".xlsx", ".docx"}
_YEAR_RE = re.compile(r"(?<!\d)((?:19|20)\d{2})(?!\d)")
_DOC_QUESTION_RE = re.compile(r"^\s*(\d+)\s*[.．、)]\s*\[([^\]]+)]\s*(.+)$")
_DOC_OPTION_RE = re.compile(r"^\s*([A-H])\s*[.．、。:]\s*(.+)$", re.IGNORECASE)
_DOC_FIELD_RE = re.compile(r"^\s*(答案|考点还原|解析)\s*[：:]\s*(.*)$")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def discover_sources(root: Path) -> list[SourceFile]:
    root = root.resolve()
    sources = []
    for path in sorted(root.rglob("*"), key=lambda item: item.as_posix().casefold()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        relative = path.relative_to(root)
        parts = relative.parts
        exam = parts[0] if len(parts) >= 2 else ""
        collection_kind = parts[1] if len(parts) >= 3 else ""
        subject = parts[2] if len(parts) >= 4 else path.parent.name
        title = path.stem
        year_match = _YEAR_RE.search("/".join(parts))
        stat = path.stat()
        sources.append(
            SourceFile(
                absolute_path=path,
                relative_path=relative.as_posix(),
                extension=path.suffix.lower(),
                sha256=sha256_file(path),
                size_bytes=stat.st_size,
                exam=display_text(exam),
                collection_kind=display_text(collection_kind),
                subject=display_text(subject),
                title=display_text(title),
                year=int(year_match.group(1)) if year_match else None,
            )
        )
    return sources


def read_questions(source: SourceFile) -> Iterator[RawQuestion]:
    if source.extension == ".xlsx":
        yield from _read_xlsx(source)
        return
    if source.extension == ".docx":
        yield from _read_docx(source)
        return
    raise ValueError(f"Unsupported source type: {source.extension}")


def _read_xlsx(source: SourceFile) -> Iterator[RawQuestion]:
    workbook = load_workbook(source.absolute_path, read_only=True, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = tuple(display_text(value) for value in next(rows, ()))
            header_index = {name: index for index, name in enumerate(headers) if name}
            if "题目" not in header_index:
                if any(headers):
                    raise ValueError(
                        f"{source.relative_path}#{worksheet.title} missing required 题目 header: {headers}"
                    )
                continue

            position = 0
            for row_number, row in enumerate(rows, start=2):
                stem = _row_value(row, header_index, "题目")
                if not stem:
                    continue
                position += 1
                options = tuple(
                    value
                    for label in "ABCDEFGH"
                    if (value := _row_value(row, header_index, f"选项{label}"))
                )
                yield RawQuestion(
                    source=source,
                    sheet_name=worksheet.title,
                    row_number=row_number,
                    position=position,
                    raw_type=_row_value(row, header_index, "题型"),
                    stem=stem,
                    options=options,
                    answer=_row_value(row, header_index, "答案"),
                    point=_row_value(row, header_index, "考点还原"),
                    explanation=_row_value(row, header_index, "答案解析"),
                )
    finally:
        workbook.close()


def _row_value(row: tuple, header_index: dict[str, int], header: str) -> str:
    index = header_index.get(header)
    if index is None or index >= len(row):
        return ""
    return display_text(row[index])


def _read_docx(source: SourceFile) -> Iterator[RawQuestion]:
    document = Document(source.absolute_path)
    current: dict | None = None
    position = 0

    def finish() -> RawQuestion | None:
        nonlocal current
        if not current or not current["stem"]:
            current = None
            return None
        question = RawQuestion(
            source=source,
            sheet_name="正文",
            row_number=current["row_number"],
            position=current["position"],
            raw_type=current["raw_type"],
            stem=current["stem"],
            options=tuple(current["options"]),
            answer=current["answer"],
            point=current["point"],
            explanation=current["explanation"],
        )
        current = None
        return question

    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        text = display_text(paragraph.text)
        if not text:
            continue
        question_match = _DOC_QUESTION_RE.match(text)
        if question_match:
            previous = finish()
            if previous:
                yield previous
            position += 1
            current = {
                "row_number": paragraph_number,
                "position": position,
                "raw_type": display_text(question_match.group(2)),
                "stem": display_text(question_match.group(3)),
                "options": [],
                "answer": "",
                "point": "",
                "explanation": "",
                "last_field": "stem",
            }
            continue
        if current is None:
            continue

        option_match = _DOC_OPTION_RE.match(text)
        if option_match:
            expected = chr(65 + len(current["options"]))
            if option_match.group(1).upper() != expected:
                raise ValueError(
                    f"{source.relative_path} paragraph {paragraph_number}: "
                    f"expected option {expected}, got {option_match.group(1).upper()}"
                )
            current["options"].append(display_text(option_match.group(2)))
            current["last_field"] = "options"
            continue

        field_match = _DOC_FIELD_RE.match(text)
        if field_match:
            field = {"答案": "answer", "考点还原": "point", "解析": "explanation"}[
                field_match.group(1)
            ]
            current[field] = display_text(field_match.group(2))
            current["last_field"] = field
            continue

        last_field = current.get("last_field")
        if last_field in {"stem", "answer", "point", "explanation"}:
            current[last_field] = display_text(f"{current[last_field]} {text}")
        elif last_field == "options" and current["options"]:
            current["options"][-1] = display_text(f"{current['options'][-1]} {text}")

    final = finish()
    if final:
        yield final
